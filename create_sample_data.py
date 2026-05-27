"""运行此文件生成测试用的 Excel 样本数据"""
import openpyxl
from pathlib import Path

DATA_DIR = Path(__file__).parent / "data" / "input"
DATA_DIR.mkdir(parents=True, exist_ok=True)

# --- sales_2024_01.xlsx ---
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "销售数据"
ws.append(["商品", "单价", "数量", "销售日期"])
ws.append(["键盘", 299, 10, "2024-01-05"])
ws.append(["鼠标", 129, 25, "2024-01-06"])
ws.append(["显示器", 1899, 3, "2024-01-10"])
ws.append(["耳机", 399, 15, "2024-01-12"])
ws.append(["鼠标垫", 39, 50, "2024-01-15"])
wb.save(DATA_DIR / "sales_2024_01.xlsx")

# --- sales_2024_02.xlsx ---
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "销售数据"
ws.append(["商品", "单价", "数量", "销售日期"])
ws.append(["键盘", 299, 8, "2024-02-03"])
ws.append(["摄像头", 499, 6, "2024-02-08"])
ws.append(["显示器", 1899, 5, "2024-02-14"])
ws.append(["鼠标", 129, 30, "2024-02-20"])
wb.save(DATA_DIR / "sales_2024_02.xlsx")

# --- sales_2024_03.xlsx ---
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "销售数据"
ws.append(["商品", "单价", "数量", "销售日期"])
ws.append(["键盘", 299, 12, "2024-03-01"])
ws.append(["耳机", 399, 20, "2024-03-05"])
ws.append(["摄像头", 499, 8, "2024-03-10"])
ws.append(["鼠标垫", 39, 100, "2024-03-15"])
ws.append(["显示器", 1899, 2, "2024-03-20"])
ws.append(["键盘", 299, 5, "2024-03-25"])
wb.save(DATA_DIR / "sales_2024_03.xlsx")

print(f"已生成 3 个样本文件到 {DATA_DIR}")
